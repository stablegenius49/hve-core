"""Validate PPTX properties that cannot be detected from rendered images.

Checks speaker notes and slide count. Visual checks (overlay, overflow,
spacing, contrast, margins, placeholders) are performed by the agent
through direct inspection of rendered slide images.

Usage:
    python validate_deck.py --input slide-deck/presentation.pptx --content-dir content/
"""

import argparse
import sys
from pathlib import Path

from pptx import Presentation

EXIT_SUCCESS = 0
EXIT_FAILURE = 1
EXIT_ERROR = 2


def check_speaker_notes(slide, slide_num: int) -> list[str]:
    """Check for missing or empty speaker notes."""
    issues = []
    try:
        notes = slide.notes_slide.notes_text_frame.text.strip()
        if not notes:
            issues.append(f"Slide {slide_num}: Empty speaker notes")
    except (AttributeError, TypeError):
        issues.append(f"Slide {slide_num}: Missing speaker notes")
    return issues


def validate_deck(pptx_path: Path, content_dir: Path | None = None,
                  slide_filter: set[int] | None = None) -> list[str]:
    """Run PPTX-only validation checks (speaker notes, slide count)."""
    prs = Presentation(str(pptx_path))
    all_issues = []

    for i, slide in enumerate(prs.slides):
        slide_num = i + 1
        if slide_filter and slide_num not in slide_filter:
            continue
        all_issues.extend(check_speaker_notes(slide, slide_num))

    if content_dir and not slide_filter:
        slide_dirs = sorted(
            [d for d in content_dir.iterdir() if d.is_dir() and d.name.startswith("slide-")]
        )
        if len(slide_dirs) == len(prs.slides):
            pass
        elif len(slide_dirs) < len(prs.slides):
            all_issues.append(
                f"Info: Partial content detected — {len(prs.slides)} slides in PPTX, "
                f"{len(slide_dirs)} content directories (expected for incremental updates)"
            )
        else:
            all_issues.append(
                f"Slide count mismatch: {len(prs.slides)} slides in PPTX, "
                f"{len(slide_dirs)} content directories"
            )

    return all_issues


def create_parser() -> argparse.ArgumentParser:
    """Create and configure argument parser."""
    parser = argparse.ArgumentParser(
        description="Validate PPTX-only properties (speaker notes, slide count)"
    )
    parser.add_argument("--input", required=True, type=Path, help="Input PPTX file path")
    parser.add_argument("--content-dir", type=Path, help="Content directory for slide count comparison")
    parser.add_argument("--slides", help="Comma-separated slide numbers to validate (default: all)")
    return parser


def main() -> int:
    """Main entry point."""
    parser = create_parser()
    args = parser.parse_args()

    pptx_path = args.input
    if not pptx_path.exists():
        print(f"Error: File not found: {pptx_path}", file=sys.stderr)
        return EXIT_ERROR

    slide_filter = None
    if args.slides:
        slide_filter = {int(s.strip()) for s in args.slides.split(",")}

    print(f"Validating PPTX properties: {pptx_path}")
    issues = validate_deck(pptx_path, args.content_dir, slide_filter=slide_filter)

    if issues:
        print(f"\n{len(issues)} issue(s) found:\n")
        for issue in issues:
            print(f"  - {issue}")
        return EXIT_FAILURE
    else:
        print("\nAll PPTX validation checks passed.")
        return EXIT_SUCCESS


if __name__ == "__main__":
    sys.exit(main())
